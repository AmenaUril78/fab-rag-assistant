"""Load documents (md, txt, pdf, docx, html), split them into chunks, and index them in Chroma.

Usage:  python -m fabrag.ingest                 (profile from FABRAG_PROFILE, default "fab")
        python -m fabrag.ingest --profile it
        python -m fabrag.ingest --profile private   (your own docs in data/private/)
"""
from __future__ import annotations

import argparse
import html
import re
from pathlib import Path
from typing import Callable, Dict, List

from langchain_chroma import Chroma
from langchain_core.documents import Document
from langchain_text_splitters import MarkdownHeaderTextSplitter, RecursiveCharacterTextSplitter

from .config import Settings, get_settings
from .embeddings import get_embeddings

DOC_ID_RE = re.compile(r"Document IDs?:\s*([A-Z0-9\-, ]+?)\s*\|")


# ---------- file loaders: every format is converted to markdown-style text ----------
def _read_md(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="ignore")


def _read_pdf(path: Path) -> str:
    from pypdf import PdfReader

    pages = [(p.extract_text() or "").strip() for p in PdfReader(str(path)).pages]
    return "\n\n".join(f"## Page {i}\n{t}" for i, t in enumerate(pages, 1) if t)


def _read_docx(path: Path) -> str:
    import docx

    out = []
    for para in docx.Document(str(path)).paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style = (para.style.name or "").lower()
        if style.startswith("title") or style == "heading 1":
            out.append(f"# {text}")
        elif style.startswith("heading"):
            out.append(f"## {text}")
        elif "list" in style:
            out.append(f"- {text}")
        else:
            out.append(text)
    return "\n\n".join(out)


def _read_html(path: Path) -> str:
    raw = path.read_text(encoding="utf-8", errors="ignore")
    raw = re.sub(r"(?is)<(script|style).*?</\1>", "", raw)
    raw = re.sub(r"(?is)<h1[^>]*>(.*?)</h1>", r"\n# \1\n", raw)
    raw = re.sub(r"(?is)<h[2-3][^>]*>(.*?)</h[2-3]>", r"\n## \1\n", raw)
    raw = re.sub(r"(?is)<li[^>]*>", "\n- ", raw)
    raw = re.sub(r"(?is)<(br|/p|/div|/tr)[^>]*>", "\n", raw)
    text = html.unescape(re.sub(r"<[^>]+>", "", raw))
    return re.sub(r"\n{3,}", "\n\n", text).strip()


LOADERS: Dict[str, Callable[[Path], str]] = {
    ".md": _read_md, ".markdown": _read_md, ".txt": _read_md,
    ".pdf": _read_pdf, ".docx": _read_docx, ".html": _read_html, ".htm": _read_html,
}


def load_documents(docs_dir: Path) -> List[Document]:
    docs = []
    for path in list_doc_files(docs_dir):
        try:
            text = LOADERS[path.suffix.lower()](path)
        except Exception as e:  # one bad file should not stop the whole index
            print(f"  skipped {path.name}: {e}")
            continue
        if not text.strip():
            continue
        # Title = first meaningful line (skipping "## Page N" markers added for PDFs), else the file name
        lines = [l.strip().lstrip("#").strip() for l in text.splitlines() if l.strip()]
        lines = [l for l in lines if not re.fullmatch(r"Page \d+", l)]
        title = lines[0] if lines and len(lines[0]) <= 100 else path.stem.replace("_", " ")
        m = DOC_ID_RE.search(text)
        docs.append(Document(page_content=text, metadata={
            "source": str(path.relative_to(docs_dir)),
            "title": title[:120],
            "doc_id": m.group(1).strip() if m else "",
        }))
    return docs


def split_documents(docs: List[Document], chunk_size: int, chunk_overlap: int) -> List[Document]:
    """Split by markdown section first (keeps procedures together), then by size."""
    header_splitter = MarkdownHeaderTextSplitter(
        headers_to_split_on=[("#", "h1"), ("##", "section")], strip_headers=False
    )
    size_splitter = RecursiveCharacterTextSplitter(chunk_size=chunk_size, chunk_overlap=chunk_overlap)

    chunks: List[Document] = []
    for doc in docs:
        n = 0
        for section in header_splitter.split_text(doc.page_content):
            sec_name = section.metadata.get("section", "Overview")
            for piece in size_splitter.split_text(section.page_content):
                # Prefix with doc title + section so each chunk carries its own context
                content = f"[{doc.metadata['title']} > {sec_name}]\n{piece}"
                meta = {**doc.metadata, "section": sec_name,
                        "chunk_id": f"{doc.metadata['source']}::{sec_name}::{n}"}
                chunks.append(Document(page_content=content, metadata=meta))
                n += 1
    return chunks


def list_doc_files(docs_dir: Path) -> List[Path]:
    """All loadable files in a folder (recursive), excluding the folder's own README."""
    if not docs_dir.exists():
        return []
    return sorted(p for p in docs_dir.rglob("*")
                  if p.is_file() and p.suffix.lower() in LOADERS and not p.name.startswith(".")
                  and not (p.name == "README.md" and p.parent == docs_dir))


def docs_fingerprint(docs_dir: Path) -> str:
    """Changes whenever a file is added, removed, or edited, so the app can re-index automatically."""
    import hashlib

    h = hashlib.sha1()
    for p in list_doc_files(docs_dir):
        st = p.stat()
        h.update(f"{p.relative_to(docs_dir)}|{st.st_size}|{int(st.st_mtime)}".encode())
    return h.hexdigest()


def index_is_current(s: Settings) -> bool:
    f = s.persist_dir / "fingerprint.txt"
    return (s.persist_dir / "chroma.sqlite3").exists() and f.exists() and \
        f.read_text().strip() == f"{s.embedding_provider}:{docs_fingerprint(s.docs_dir)}"


def build_index(settings: Settings | None = None, reset: bool = True) -> Chroma:
    s = settings or get_settings()
    s.persist_dir.mkdir(parents=True, exist_ok=True)
    embeddings = get_embeddings(s.embedding_provider, s.openai_embedding_model)
    kwargs = dict(collection_name=s.collection, embedding_function=embeddings,
                  persist_directory=str(s.persist_dir))
    if reset:  # drop the old collection through Chroma (safe even while the app has it open)
        Chroma(**kwargs).delete_collection()
    store = Chroma(**kwargs, collection_metadata={"hnsw:space": "cosine",
                                                  "embedding_provider": s.embedding_provider})

    docs = load_documents(s.docs_dir)
    chunks = split_documents(docs, s.chunk_size, s.chunk_overlap)
    if chunks:
        store.add_documents(chunks, ids=[c.metadata["chunk_id"] for c in chunks])
    (s.persist_dir / "fingerprint.txt").write_text(f"{s.embedding_provider}:{docs_fingerprint(s.docs_dir)}")
    print(f"[{s.profile.key}] Indexed {len(chunks)} chunks from {len(docs)} documents in {s.docs_dir} "
          f"(embeddings: {s.embedding_provider})")
    return store


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--profile", default=None, help="fab | it | private")
    build_index(get_settings(ap.parse_args().profile))
