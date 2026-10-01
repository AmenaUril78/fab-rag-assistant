"""Offline tests (no API keys, no model downloads): hashing embeddings + extractive mode."""
import subprocess
from pathlib import Path

import pytest

from fabrag.config import ROOT, get_settings
from fabrag.ingest import build_index, load_documents, split_documents
from fabrag.rag import FabRAG, build_prompt, rrf_fuse, tokenize


def make_rag(profile, tmp_path_factory, **kw):
    s = get_settings(profile, persist_dir=tmp_path_factory.mktemp(f"chroma_{profile}"),
                     embedding_provider="hashing", llm_provider="none", **kw)
    build_index(s)
    return FabRAG(s)


@pytest.fixture(scope="module")
def rag(tmp_path_factory):
    return make_rag("fab", tmp_path_factory)


@pytest.fixture(scope="module")
def it_rag(tmp_path_factory):
    return make_rag("it", tmp_path_factory)


@pytest.mark.parametrize("profile", ["fab", "it"])
def test_chunks_keep_section_metadata(profile):
    s = get_settings(profile)
    chunks = split_documents(load_documents(s.docs_dir), 700, 100)
    assert len(chunks) > 20
    assert all(c.metadata["section"] and c.metadata["source"] for c in chunks)
    assert len({c.metadata["chunk_id"] for c in chunks}) == len(chunks)


def test_tokenizer_keeps_codes():
    toks = tokenize("Tool shows ALM-4021 and %PM-4-ERR_DISABLE again")
    assert "alm-4021" in toks and "4021" in toks
    assert "pm-4-err_disable" in toks and "disable" in toks


def test_rrf_prefers_docs_ranked_high_in_both_lists(rag):
    a, b, c = rag.all_chunks[:3]
    fused = rrf_fuse([[a, b, c], [b, a, c]])
    assert {fused[0][0].metadata["chunk_id"], fused[1][0].metadata["chunk_id"]} == {
        a.metadata["chunk_id"], b.metadata["chunk_id"]}


@pytest.mark.parametrize("q,expected", [
    ("What do I do for ALM-4021?", "alarm_code_reference.md"),
    ("when should the polishing pad be replaced", "cmp_process_guide.md"),
    ("gas leak response", "safety_lockout_gas.md"),
])
def test_fab_retrieval_finds_right_doc(rag, q, expected):
    assert expected in [d.metadata["source"] for d, _ in rag.retrieve(q, k=3)]


@pytest.mark.parametrize("q,expected", [
    ("Switch log shows %PM-4-ERR_DISABLE", "switch_port_troubleshooting.md"),
    ("client got a 169.254 address", "dhcp_wireless_runbook.md"),
    ("backup failed with AuthenticationException", "config_backup_runbook.md"),
])
def test_it_retrieval_finds_right_doc(it_rag, q, expected):
    assert expected in [d.metadata["source"] for d, _ in it_rag.retrieve(q, k=3)]


def test_prompt_is_profile_specific():
    fab = build_prompt(get_settings("fab")).messages[0].prompt.template
    it = build_prompt(get_settings("it")).messages[0].prompt.template
    assert "FabAssist" in fab and "safety" in fab
    assert "OpsAssist" in it and "rollback" in it


def test_ask_returns_cited_sources(rag):
    ans = rag.ask("How often do we change the SC1 bath?")
    assert ans.sources and ans.sources[0]["n"] == 1
    assert "[1]" in ans.answer


def test_loads_txt_docx_html(tmp_path):
    import docx

    (tmp_path / "notes.txt").write_text("VPN error 809 means UDP 500 is blocked by the firewall.")
    (tmp_path / "page.html").write_text("<h1>Printer Runbook</h1><h2>Jam</h2><p>Open tray 2 &amp; remove paper.</p>")
    d = docx.Document()
    d.add_heading("Laptop Imaging", level=1)
    d.add_heading("BitLocker", level=2)
    d.add_paragraph("Suspend BitLocker before a BIOS update.")
    d.save(tmp_path / "imaging.docx")
    (tmp_path / "sub").mkdir()
    (tmp_path / "sub" / "deep.md").write_text("# Deep\n\n## Part\nnested folder works")

    docs = {doc.metadata["source"]: doc for doc in load_documents(tmp_path)}
    assert set(docs) == {"notes.txt", "page.html", "imaging.docx", "sub/deep.md"}
    assert "Open tray 2 & remove paper." in docs["page.html"].page_content
    assert docs["imaging.docx"].metadata["title"] == "Laptop Imaging"
    sections = {c.metadata["section"] for c in split_documents(list(docs.values()), 700, 100)}
    assert {"BitLocker", "Jam", "Part"} <= sections


def test_private_docs_are_git_ignored():
    probe = ROOT / "data" / "private" / "secret_runbook.md"
    out = subprocess.run(["git", "check-ignore", str(probe)], cwd=ROOT, capture_output=True, text=True)
    assert out.returncode == 0, "data/private/ must be git-ignored"


def test_rebuild_index_while_open(tmp_path_factory):
    s = get_settings("it", persist_dir=tmp_path_factory.mktemp("rebuild"), embedding_provider="hashing",
                     llm_provider="none")
    build_index(s)
    first = FabRAG(s)
    build_index(s)  # rebuild while a client is still open
    assert len(FabRAG(s).all_chunks) == len(first.all_chunks)


def test_api_endpoints(rag, monkeypatch):
    from fastapi.testclient import TestClient

    import api

    monkeypatch.setattr(api, "get_rag", lambda: rag)
    client = TestClient(api.app)
    assert client.get("/health").json()["status"] == "ok"
    r = client.post("/ask", json={"question": "What is CVD-2210?", "top_k": 3})
    assert r.status_code == 200 and len(r.json()["sources"]) == 3
    r = client.post("/search", json={"query": "overlay", "mode": "bm25"})
    assert r.status_code == 200 and r.json()
